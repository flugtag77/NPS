#!/usr/bin/env python3
"""Строит витрину ЖК детально: n накоп, NPS, распределение 0-10, П/Н/К, дневные срезы."""
import openpyxl, json, re, sys, glob, os
from collections import defaultdict, Counter
from datetime import datetime

def latest_xlsx():
    cands = sorted(glob.glob('/home/user/workspace/uploaded_attachments/*/NPS-2026.xlsx'), key=os.path.getmtime, reverse=True)
    return cands[0] if cands else None

XLSX = sys.argv[1] if len(sys.argv)>1 else latest_xlsx()

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb.active

meta = json.load(open('/home/user/workspace/nps_pipeline/db/zhk_meta.json',encoding='utf-8'))
year_matrix = {}
import csv
with open('/home/user/workspace/nps_pipeline/db/zhk_year_matrix_full.csv',encoding='utf-8') as f:
    r = csv.DictReader(f)
    for row in r:
        year_matrix[row['ЖК']] = row

def norm_name(s): return re.sub(r'\s+','',str(s).lower().replace('ё','е').replace('«','').replace('»','').replace('"','').strip())
meta_norm = {norm_name(k):(k,v) for k,v in meta.items()}
ym_norm = {norm_name(k):v for k,v in year_matrix.items()}

def seg(nps):
    if nps is None or nps=='': return None
    try: v = int(nps)
    except: return None
    if v <= 6: return 'C'
    if v <= 8: return 'N'
    return 'P'

def parse_date(s):
    if not s: return None
    if isinstance(s, datetime): return s.strftime('%d.%m.%Y')
    s = str(s).strip()
    # DD.MM.YYYY HH:MM:SS
    m = re.match(r'(\d{1,2})\.(\d{1,2})\.(\d{4})', s)
    if m: return f'{int(m.group(1)):02d}.{int(m.group(2)):02d}.{m.group(3)}'
    return None

# Сбор
zhk_raw = defaultdict(list)  # zhk -> list[(nps, seg, date)]
for r in range(2, ws.max_row+1):
    nps_v = ws.cell(r, 42).value  # IDX 41 + 1
    zhk = ws.cell(r, 9).value  # IDX 8 + 1
    date = ws.cell(r, 2).value  # IDX 1 + 1
    if nps_v is None or nps_v == '': continue
    s = seg(nps_v)
    if not s: continue
    zhk_raw[str(zhk).strip() if zhk else '?'].append((int(nps_v), s, parse_date(date)))

# получаем список дат которые уже прошли
all_dates = set()
for lst in zhk_raw.values():
    for _, _, d in lst:
        if d: all_dates.add(d)
def dt_key(d):
    try:
        return datetime.strptime(d, '%d.%m.%Y')
    except: return datetime(2000,1,1)
sorted_dates = sorted(all_dates, key=dt_key)
print('dates:', sorted_dates)

# Периметр 96
perimeter = []
with open('/home/user/workspace/nps_pipeline/db/zhk_full.tsv',encoding='utf-8') as f:
    r = csv.reader(f, delimiter='\t')
    next(r)
    for row in r:
        perimeter.append(row[1])

rows = []
for z in perimeter:
    n = norm_name(z)
    m = meta.get(z, {})
    ym = ym_norm.get(n, {})
    entries = zhk_raw.get(z, [])
    # ещё пробуем матчить по нормализации
    if not entries:
        for raw_key, lst in zhk_raw.items():
            if norm_name(raw_key) == n:
                entries = lst; break
    total_n = len(entries)
    c = sum(1 for _,s,_ in entries if s=='C')
    nn = sum(1 for _,s,_ in entries if s=='N')
    p = sum(1 for _,s,_ in entries if s=='P')
    nps_v = round((p-c)/total_n*100,1) if total_n else None
    # распределение 0-10 в %
    dist_cnt = Counter(x[0] for x in entries)
    dist_pct = {}
    for i in range(11):
        dist_pct[i] = round(dist_cnt[i]/total_n*100,1) if total_n else 0
    # дневные срезы
    daily = {}
    for d in sorted_dates:
        day_ent = [e for e in entries if e[2]==d]
        n_d = len(day_ent)
        if n_d:
            cd = sum(1 for _,s,_ in day_ent if s=='C')
            pd = sum(1 for _,s,_ in day_ent if s=='P')
            npsd = round((pd-cd)/n_d*100,1)
        else:
            npsd = None
        daily[d] = {'n': n_d, 'nps': npsd}
    rows.append({
        'name': z,
        'region': m.get('region',''),
        'klass': m.get('klass',''),
        'event': m.get('event',''),
        'n': total_n,
        'nps': nps_v,
        'pct_c': round(c/total_n*100,1) if total_n else 0,
        'pct_n': round(nn/total_n*100,1) if total_n else 0,
        'pct_p': round(p/total_n*100,1) if total_n else 0,
        'dist': dist_pct,
        'nps_2025': ym.get('NPS_2025',''),
        'n_2025': ym.get('n_2025',''),
        'nps_2024': ym.get('NPS_2024',''),
        'n_2024': ym.get('n_2024',''),
        'daily': daily,
    })

rows.sort(key=lambda x: -x['n'])
out = {'dates': sorted_dates, 'rows': rows}
with open('/home/user/workspace/nps_pipeline/db/zhk_detail.json','w',encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
print('saved zhk_detail.json — rows:', len(rows))
print('sample:', rows[0]['name'], 'n=', rows[0]['n'], 'nps=', rows[0]['nps'])
