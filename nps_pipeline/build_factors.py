#!/usr/bin/env python3
"""Считает все факторы влияния на NPS из xlsx выгрузки."""
import openpyxl, json, re, sys, glob, os
from collections import defaultdict, Counter

def latest_xlsx():
    cands = sorted(glob.glob('/home/user/workspace/uploaded_attachments/*/NPS-2026.xlsx'), key=os.path.getmtime, reverse=True)
    return cands[0] if cands else None

XLSX = sys.argv[1] if len(sys.argv)>1 else latest_xlsx()
print('XLSX:', XLSX)

IDX = {"DATE":1, "ZHK":8, "TENURE":9, "GENDER":10, "AGE":11,
       "OFFICE_VISIT":12, "OFFICE_FREQ":13, "NPS":41,
       "GARANT_ASK":42, "GARANT_RATE":43, "GARANT_IMPACT":44}

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb.active

meta = json.load(open('/home/user/workspace/nps_pipeline/db/zhk_meta.json',encoding='utf-8'))
def norm(s): return re.sub(r'\s+','',str(s).lower().replace('ё','е').replace('«','').replace('»','').replace('"','').strip())
meta_norm = {norm(k):v for k,v in meta.items()}

def seg(nps):
    if nps is None or nps=='': return None
    try: v = int(nps)
    except: return None
    if v <= 6: return 'C'
    if v <= 8: return 'N'
    return 'P'

def nps_of(rows):
    n = len(rows)
    if n == 0: return {'n':0,'nps':None,'pct_c':0,'pct_n':0,'pct_p':0,'c':0,'nn':0,'p':0}
    c = sum(1 for r in rows if r=='C')
    p = sum(1 for r in rows if r=='P')
    nn = sum(1 for r in rows if r=='N')
    return {'n':n,'nps':round((p-c)/n*100,1),'pct_c':round(c/n*100,1),'pct_n':round(nn/n*100,1),'pct_p':round(p/n*100,1),'c':c,'nn':nn,'p':p}

# --- нормализаторы значений ---
def norm_office_visit(v):
    if not v: return None
    v = str(v).strip()
    if v.startswith('Да'): return 'Да'
    if 'Нет' in v: return 'Нет'
    return None

def norm_age(v):
    """31-40 лет -> '31–40'"""
    if not v: return None
    v = str(v).strip().lower()
    if 'до 20' in v or 'младше 20' in v or '<20' in v or 'моложе 20' in v: return '< 20'
    m = re.search(r'(\d+)\s*[-–]\s*(\d+)', v)
    if m:
        a,b = int(m.group(1)), int(m.group(2))
        if b<=20: return '< 20'
        if b<=30: return '21–30'
        if b<=40: return '31–40'
        if b<=50: return '41–50'
        return '> 50'
    if 'старше 50' in v or 'больше 50' in v or '50+' in v or '>50' in v: return '> 50'
    return None

def norm_impact(v):
    if not v: return None
    v = str(v).strip().lower()
    if 'положительн' in v or 'улучшил' in v: return 'Улучшил моё отношение'
    if 'никак' in v or 'не повлия' in v: return 'Не повлиял'
    if 'отрицательн' in v or 'ухудшил' in v: return 'Ухудшил моё отношение'
    return None

def norm_freq(v):
    if not v: return None
    v = str(v).strip().lower()
    if 'нескольк' in v and 'месяц' in v: return 'Несколько раз в месяц'
    if '2' in v and ('раза' in v or 'раз ' in v): return '2–3 раза за год'
    if 'один раз' in v or '1 раз' in v: return '1 раз за год'
    if 'реже' in v: return 'Реже раза в год'
    if 'первы' in v: return 'Первый раз'
    return str(v).strip().capitalize()

data = []
for r in range(2, ws.max_row+1):
    row = [ws.cell(r,c+1).value for c in range(50)]
    s = seg(row[IDX['NPS']])
    if not s: continue
    zhk = str(row[IDX['ZHK']]).strip() if row[IDX['ZHK']] else ''
    m = meta_norm.get(norm(zhk), {})
    data.append({
        'seg': s,
        'gender': str(row[IDX['GENDER']] or '').strip() or None,
        'age': norm_age(row[IDX['AGE']]),
        'tenure': str(row[IDX['TENURE']] or '').strip() or None,
        'office_visit': norm_office_visit(row[IDX['OFFICE_VISIT']]),
        'office_freq': norm_freq(row[IDX['OFFICE_FREQ']]),
        'garant_ask': str(row[IDX['GARANT_ASK']] or '').strip() or None,
        'garant_rate': row[IDX['GARANT_RATE']],
        'garant_impact': norm_impact(row[IDX['GARANT_IMPACT']]),
        'region': m.get('region',''),
        'event': m.get('event',''),
    })
print('total:', len(data))

def bucket(field, order=None):
    buckets = defaultdict(list)
    for d in data:
        v = d[field]
        if v: buckets[v].append(d['seg'])
    keys = order if order else sorted(buckets.keys())
    out = []
    for k in keys:
        if k in buckets:
            r = nps_of(buckets[k]); r['name']=k
            out.append(r)
    # добавить те что вне order
    if order:
        for k in buckets:
            if k not in order:
                r = nps_of(buckets[k]); r['name']=k
                out.append(r)
    return out

factors = {
    'gender': bucket('gender', ['Мужской','Женский']),
    'age': bucket('age', ['< 20','21–30','31–40','41–50','> 50']),
    'region': bucket('region', ['Москва','МО','Санкт-Петербург','ЛО','Москва БК Продукт','Нижегородская обл.','Свердловская обл.','Татарстан']),
    'tenure': bucket('tenure', ['Нахожусь на этапе заселения','Менее 1 месяца назад','От 1 до 6 месяцев назад','От 6 месяцев до 1 года назад','Более 1 года назад']),
    'office_visit': bucket('office_visit', ['Да','Нет']),
    'office_freq': bucket('office_freq'),
    'garant_ask': bucket('garant_ask', ['Да','Нет']),
    'garant_impact': bucket('garant_impact', ['Улучшил моё отношение','Не повлиял','Ухудшил моё отношение']),
    'event': bucket('event', ['Группа А (программа максимум)','Группа B (программа минимум)','Группа С (не было мероприятий)']),
}

# распределение оценки гарантии 1-5
rate_dist = Counter()
total_rate = 0
for d in data:
    if d['garant_ask']=='Да' and d['garant_rate'] is not None:
        try:
            v = int(d['garant_rate'])
            if 1<=v<=5: rate_dist[v]+=1; total_rate+=1
        except: pass
factors['garant_rate'] = {
    'total': total_rate,
    'dist': {str(k): rate_dist[k] for k in [1,2,3,4,5]},
    'dist_pct': {str(k): round(rate_dist[k]/max(total_rate,1)*100,1) for k in [1,2,3,4,5]},
    'avg': round(sum(k*v for k,v in rate_dist.items())/max(total_rate,1),2),
    'csat45': round((rate_dist[4]+rate_dist[5])/max(total_rate,1)*100,1),
}

# сохраняю
out = '/home/user/workspace/nps_pipeline/db/factors.json'
with open(out,'w',encoding='utf-8') as f:
    json.dump(factors, f, ensure_ascii=False, indent=1)
print('saved', out)
for k,v in factors.items():
    if isinstance(v, list):
        print(f'  {k}: {len(v)} groups, total n={sum(x["n"] for x in v)}')
    else:
        print(f'  {k}: total={v.get("total",0)}')
