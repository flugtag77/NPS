"""Строит baseline 2024/2025 из xlsx-выгрузок Testograf.

Читает NPS-2025.xlsx и YGC-2024.xlsx, матчит ЖК к периметру 2026,
считает по каждому году: n, NPS, CSAT, промо/нейтр/крит,
+ строит кривую набора по дням для калибровки прогноза объёма 2026.

Перезаписывает db/zhk_year_matrix_full.csv (реальные NPS_2024, NPS_2025).
Сохраняет db/curve_2025.json — % набора по дням от старта опроса.
Никаких LLM, только Python.
"""
import csv, json, re, os
from openpyxl import load_workbook
from collections import defaultdict
from datetime import datetime

ROOT = "/home/user/workspace/nps_pipeline"

PATH_2025 = "/home/user/workspace/uploaded_attachments/34dcbdfb42ac44f69753cae3d0526b17/NPS-2025.xlsx"
PATH_2024 = "/home/user/workspace/uploaded_attachments/34dcbdfb42ac44f69753cae3d0526b17/YGC-2024.xlsx"

# Колонки (0-based)
COL_2025 = {"date": 1, "zhk": 8, "csat": 36, "nps": 65}
COL_2024 = {"date": 1, "zhk": 8, "csat": 36, "nps": 64}

# --- Периметр 2026 ---
perim = []
with open(f"{ROOT}/db/zhk_full.tsv", encoding="utf-8") as f:
    for r in csv.DictReader(f, delimiter="\t"):
        perim.append(r["zhk"])
perim_set = set(perim)

# --- Нормализация имён ЖК ---
def norm(s):
    if not s: return ""
    s = str(s).strip().lower()
    s = s.replace("ё", "е")
    s = re.sub(r"[«»\"']", "", s)
    s = re.sub(r"\s+", " ", s)
    return s

perim_norm = {norm(z): z for z in perim}

def match(raw):
    """Возвращает канон-имя ЖК из периметра 2026 или None."""
    if not raw: return None
    n = norm(raw)
    if n in perim_norm:
        return perim_norm[n]
    # частичное совпадение — если normalized(perim) полностью содержится в normalized(raw) или наоборот
    for pn, canon in perim_norm.items():
        if pn == n: return canon
        if len(pn) >= 5 and (pn in n or n in pn):
            return canon
    return None

# --- Парсим один xlsx ---
def parse(path, col, year):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(min_row=2, values_only=True)
    agg = defaultdict(lambda: {"n":0, "p":0, "neu":0, "d":0, "csat_sum":0, "csat_n":0, "nps_sum":0})
    daily = defaultdict(int)  # date -> count
    total = 0
    unmatched = defaultdict(int)
    for row in it:
        if not row or len(row) <= col["nps"]: continue
        zhk_raw = row[col["zhk"]]
        nps_v = row[col["nps"]]
        if nps_v is None: continue
        try: nps = int(float(nps_v))
        except: continue
        if not (0 <= nps <= 10): continue
        z = match(zhk_raw)
        if z is None:
            unmatched[str(zhk_raw)[:60]] += 1
            continue
        total += 1
        a = agg[z]
        a["n"] += 1
        a["nps_sum"] += nps
        if nps >= 9: a["p"] += 1
        elif nps >= 7: a["neu"] += 1
        else: a["d"] += 1
        # CSAT
        try:
            c = float(row[col["csat"]])
            if 1 <= c <= 5:
                a["csat_sum"] += c
                a["csat_n"] += 1
        except: pass
        # Дневная динамика
        d = row[col["date"]]
        d_iso = None
        if isinstance(d, datetime):
            d_iso = d.date().isoformat()
        elif isinstance(d, str) and len(d) >= 10:
            s = d[:10]
            # DD.MM.YYYY → YYYY-MM-DD
            if s[2] == "." and s[5] == ".":
                d_iso = f"{s[6:10]}-{s[3:5]}-{s[0:2]}"
            elif s[4] == "-":
                d_iso = s
        if d_iso:
            daily[d_iso] += 1
    wb.close()
    return agg, daily, total, unmatched

print("[1/3] Парсинг 2024...")
agg24, daily24, tot24, unm24 = parse(PATH_2024, COL_2024, 2024)
print(f"  Всего валидных ответов 2024: {tot24}, матчено ЖК: {len(agg24)}")
if unm24:
    top_unm = sorted(unm24.items(), key=lambda x:-x[1])[:5]
    print(f"  ТОП не-матч 2024: {top_unm}")

print("[2/3] Парсинг 2025...")
agg25, daily25, tot25, unm25 = parse(PATH_2025, COL_2025, 2025)
print(f"  Всего валидных ответов 2025: {tot25}, матчено ЖК: {len(agg25)}")
if unm25:
    top_unm = sorted(unm25.items(), key=lambda x:-x[1])[:5]
    print(f"  ТОП не-матч 2025: {top_unm}")

# --- 2026 факт из responses_long.csv (только матчинг с периметром, для полноты таблицы) ---
agg26 = defaultdict(lambda: {"n":0,"p":0,"neu":0,"d":0})
with open(f"{ROOT}/db/responses_long.csv", encoding="utf-8") as f:
    for r in csv.DictReader(f):
        if r["year"] != "2026": continue
        z = r["zhk"]
        if z not in perim_set: continue
        try: nps = int(r["nps_uk"])
        except: continue
        a = agg26[z]
        a["n"] += 1
        if nps >= 9: a["p"] += 1
        elif nps >= 7: a["neu"] += 1
        else: a["d"] += 1

# --- Функции ---
def nps_of(a):
    if a["n"] == 0: return None
    return round((a["p"] - a["d"]) / a["n"] * 100, 1)

def hyp(n):
    return "" if n and n >= 50 else "гипотеза"

# --- Запись матрицы ---
print("[3/3] Запись zhk_year_matrix_full.csv...")
out_path = f"{ROOT}/db/zhk_year_matrix_full.csv"
# Бэкап старой
if os.path.exists(out_path):
    os.rename(out_path, out_path + ".bak")

with open(out_path, "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f)
    w.writerow(["ЖК","NPS_2024","n_2024","CSAT_2024","NPS_2025","n_2025","CSAT_2025",
                "NPS_2026","n_2026","Δ 2025 vs 2024","hypothesis_2025_2024",
                "Δ 2026 vs 2025","hypothesis_2026_2025",
                "Δ 2026 vs 2024","hypothesis_2026_2024"])
    for z in perim:
        a24 = agg24.get(z, {"n":0,"p":0,"d":0,"csat_sum":0,"csat_n":0})
        a25 = agg25.get(z, {"n":0,"p":0,"d":0,"csat_sum":0,"csat_n":0})
        a26 = agg26.get(z, {"n":0,"p":0,"d":0})
        n24, n25, n26 = a24["n"], a25["n"], a26["n"]
        nps24 = nps_of(a24)
        nps25 = nps_of(a25)
        nps26 = nps_of(a26)
        csat24 = round(a24["csat_sum"]/a24["csat_n"],2) if a24["csat_n"] else ""
        csat25 = round(a25["csat_sum"]/a25["csat_n"],2) if a25["csat_n"] else ""
        d_25_24 = round(nps25 - nps24, 1) if (nps25 is not None and nps24 is not None) else ""
        d_26_25 = round(nps26 - nps25, 1) if (nps26 is not None and nps25 is not None) else ""
        d_26_24 = round(nps26 - nps24, 1) if (nps26 is not None and nps24 is not None) else ""
        w.writerow([z,
                    nps24 if nps24 is not None else "", n24, csat24,
                    nps25 if nps25 is not None else "", n25, csat25,
                    nps26 if nps26 is not None else "", n26,
                    d_25_24, hyp(min(n24,n25) if n24 and n25 else 0),
                    d_26_25, hyp(min(n26,n25) if n26 and n25 else 0),
                    d_26_24, hyp(min(n26,n24) if n26 and n24 else 0)])

# --- Кривая набора 2025 (от старта опроса) ---
if daily25:
    sorted_days = sorted(daily25.items())
    start = datetime.fromisoformat(sorted_days[0][0]).date()
    cumul = 0
    total_2025 = sum(daily25.values())
    curve = []
    for date_s, cnt in sorted_days:
        d = datetime.fromisoformat(date_s).date()
        day_idx = (d - start).days + 1
        cumul += cnt
        curve.append({"day": day_idx, "date": date_s, "count": cnt,
                      "cumul": cumul, "pct": round(cumul/total_2025*100, 2)})
    with open(f"{ROOT}/db/curve_2025.json", "w", encoding="utf-8") as f:
        json.dump({"total": total_2025, "days": len(curve), "curve": curve}, f, ensure_ascii=False, indent=2)
    print(f"  Кривая 2025: {len(curve)} дней, всего {total_2025} ответов")
    # Ключевые точки
    for target_day in [1, 3, 4, 5, 7, 10, 14, 21, 24]:
        pt = next((x for x in curve if x["day"] == target_day), None)
        if pt: print(f"    День {target_day}: {pt['cumul']} ({pt['pct']}%)")

# --- Итоговые агрегаты ---
def overall_nps(agg):
    tp = sum(a["p"] for a in agg.values())
    td = sum(a["d"] for a in agg.values())
    tn = sum(a["n"] for a in agg.values())
    return round((tp-td)/tn*100, 1) if tn else 0, tn

nps24_all, n24_all = overall_nps(agg24)
nps25_all, n25_all = overall_nps(agg25)
nps26_all, n26_all = overall_nps(agg26)
print()
print(f"=== ИТОГИ ПО ПЕРИМЕТРУ 2026 ===")
print(f"2024: NPS = {nps24_all}  |  n = {n24_all}")
print(f"2025: NPS = {nps25_all}  |  n = {n25_all}")
print(f"2026: NPS = {nps26_all}  |  n = {n26_all}")
print()
print(f"Матрица: {out_path}")
