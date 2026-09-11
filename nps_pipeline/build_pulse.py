#!/usr/bin/env python3
"""Пульс_день: накопительный NPS, распределение по дням, коммуникация."""
import openpyxl, json, sys, glob, os, re
from collections import defaultdict, Counter
from datetime import datetime

def latest_xlsx():
    return sorted(glob.glob('/home/user/workspace/uploaded_attachments/*/NPS-2026.xlsx'), key=os.path.getmtime, reverse=True)[0]

XLSX = sys.argv[1] if len(sys.argv)>1 else latest_xlsx()
wb = openpyxl.load_workbook(XLSX, data_only=True); ws = wb.active

def seg(v):
    if v is None or v=='': return None
    try: v=int(v)
    except: return None
    if v<=6: return 'C'
    if v<=8: return 'N'
    return 'P'

def parse_date(s):
    if not s: return None
    if isinstance(s, datetime): return s.strftime('%d.%m.%Y')
    m = re.match(r'(\d{1,2})\.(\d{1,2})\.(\d{4})', str(s).strip())
    if m: return f'{int(m.group(1)):02d}.{int(m.group(2)):02d}.{m.group(3)}'
    return None

# сбор день -> [(nps, seg)]
day_data = defaultdict(list)
for r in range(2, ws.max_row+1):
    v = ws.cell(r,42).value
    s = seg(v)
    if not s: continue
    d = parse_date(ws.cell(r,2).value)
    if d: day_data[d].append((int(v), s))

def dt_key(d):
    try: return datetime.strptime(d, '%d.%m.%Y')
    except: return datetime(2000,1,1)
dates = sorted(day_data.keys(), key=dt_key)

# коммуникации
COMMS = {
    '07.09.2026': 'Рассылка в MAX и TG каналах',
    '08.09.2026': 'Пост-МиП с праздников с ссылкой на NPS',
    '09.09.2026': 'PUSH в ЛКК: Петр 1, Лефортово парк, Ясеневая, Люберцы парк, Митино парк',
    '10.09.2026': 'PUSH в ЛКК: Второй Нагатинский, Ильинские луга, Белая дача парк, Жемчужина Зеленограда, Академика Павлова, Открытый парк, Одинцово 1, Никольские луга, Черняховского 19, Зеленый парк',
    '11.09.2026': 'PUSH в ЛК: Красноказарменная 15, Столичные поляны, Середневский лес, Green park, Измайловский лес, Мякинино парк, Новохохловская 15, Барклая 6, Кузьминский лес, Кольская 8',
}

rows = []
cum = []
for d in dates:
    day = day_data[d]
    cum.extend(day)
    n_cum = len(cum); n_day = len(day)
    c_c = sum(1 for _,s in cum if s=='C')
    n_c = sum(1 for _,s in cum if s=='N')
    p_c = sum(1 for _,s in cum if s=='P')
    nps_c = round((p_c-c_c)/n_cum*100,1) if n_cum else None
    c_d = sum(1 for _,s in day if s=='C')
    n_d = sum(1 for _,s in day if s=='N')
    p_d = sum(1 for _,s in day if s=='P')
    nps_d = round((p_d-c_d)/n_day*100,1) if n_day else None
    dist_cnt = Counter(v for v,_ in cum)
    dist_pct = {i: round(dist_cnt[i]/n_cum*100,1) if n_cum else 0 for i in range(11)}
    rows.append({
        'date': d,
        'n_cum': n_cum, 'n_day': n_day,
        'nps_cum': nps_c, 'nps_day': nps_d,
        'crit_cum': c_c, 'neut_cum': n_c, 'promo_cum': p_c,
        'pct_c_cum': round(c_c/n_cum*100,1) if n_cum else 0,
        'pct_n_cum': round(n_c/n_cum*100,1) if n_cum else 0,
        'pct_p_cum': round(p_c/n_cum*100,1) if n_cum else 0,
        'pct_c_day': round(c_d/n_day*100,1) if n_day else 0,
        'pct_n_day': round(n_d/n_day*100,1) if n_day else 0,
        'pct_p_day': round(p_d/n_day*100,1) if n_day else 0,
        'dist_cum': dist_pct,
        'comm': COMMS.get(d, ''),
    })

with open('/home/user/workspace/nps_pipeline/db/pulse.json','w',encoding='utf-8') as f:
    json.dump({'rows':rows}, f, ensure_ascii=False, indent=1)
print('saved pulse.json — days:', len(rows))
for r in rows:
    print(f'  {r["date"]}: n={r["n_cum"]} (+{r["n_day"]}) NPS={r["nps_cum"]}% day={r["nps_day"]}%')
