"""
recompute_all.py — пересчёт всей аналитики из raw xlsx 68e94bf8.
Генерирует свежие SUMMARY, SHEETS, MAP_ZHK, FORECAST_VOL, STAB_TOTAL.
Вставляет их в _snapshots/09_46_v2.html.
"""
import json, re, datetime as dt
from collections import Counter, defaultdict
from openpyxl import load_workbook

import sys, glob
# Аргумент --xlsx имеет приоритет, иначе берём самый свежий NPS-2026.xlsx в uploaded_attachments
if "--xlsx" in sys.argv:
    SRC = sys.argv[sys.argv.index("--xlsx") + 1]
else:
    _candidates = glob.glob("/home/user/workspace/uploaded_attachments/*/NPS-2026.xlsx")
    SRC = max(_candidates, key=lambda p: __import__("os").path.getmtime(p)) if _candidates else "/home/user/workspace/uploaded_attachments/68e94bf81cf847678ecb84bf37c419ce/NPS-2026.xlsx"
print(f"Использую xlsx: {SRC}")
SNAP = "/home/user/workspace/nps_pipeline/_snapshots/09_46_v2.html"
PERIMETER = "/home/user/workspace/nps_pipeline/db/zhk_full.tsv"

# Периметр 96 ЖК
perimeter = []
import csv
with open(PERIMETER, encoding="utf-8") as f:
    for r in csv.DictReader(f, delimiter="\t"):
        perimeter.append({"zhk": r["zhk"], "reg": r["reg"], "meropr": r["meropr"]})
PERIM_SET = {p["zhk"] for p in perimeter}
print(f"Периметр: {len(perimeter)} ЖК")

# 2025 факты по ЖК — читаю из старого перимера
zhk_2025 = {}
with open("/home/user/workspace/nps_pipeline/db/zhk_year_matrix_full.csv", encoding="utf-8") as f:
    for r in csv.DictReader(f):
        z = r.get("zhk") or r.get("ЖК")
        if z:
            def _f(k):
                v = r.get(k, "")
                try: return float(v) if v not in ("","None",None) else None
                except: return None
            def _i(k):
                v = r.get(k, "")
                try: return int(float(v)) if v not in ("","None",None) else 0
                except: return 0
            zhk_2025[z] = {"n_2025": _i("n_2025"), "nps_2025": _f("nps_2025")}
print(f"ЖК 2025 факт: {len(zhk_2025)}")

# --- Читаю raw xlsx ---
wb = load_workbook(SRC, data_only=True)
ws = wb["Sheet1"]
rows = list(ws.iter_rows(values_only=True))
header = rows[0]
data = rows[1:]

IDX = {
    "ID": 0, "DATE": 1, "OS": 5, "DEV": 7, "ZHK": 8,
    "TENURE": 9, "GENDER": 10, "AGE": 11, "OFFICE_VISIT": 12, "OFFICE_FREQ": 13,
    "CSAT": 14, "NPS": 41, "GARANT_ASK": 42, "GARANT_RATE": 43, "GARANT_IMPACT": 44,
    "OPYT": 45, "OPEN": 46,
    # CSI критерии — каждая пара: [i]=оценка, [i+1]=важность
    "CSI_START": 15,  # до 40, шаг 2, 13 критериев
}
CSI_LABELS = [
    "Уборка мест общего пользования",
    "Уборка и содержание территории",
    "Обслуживание улич. освещения",
    "Качество озеленения",
    "Работа лифтового оборудования",
    "Работа пожарной сигнализации",
    "Работа инженерных систем",
    "Оперативность выездов",
    "Информирование",
    "Организация состояния подъездов",
    "Работа системы контроля доступа",
    "Доступность и эффективность оф.",
    "Удобство и функционал ЛКК/приложения",
]

def _int(v, lo=0, hi=10):
    if v is None or v == "": return None
    try:
        iv = int(v)
        return iv if lo <= iv <= hi else None
    except: return None

# --- Парсим все ответы ---
records = []
for r in data:
    if r[IDX["ID"]] is None: continue
    rec = {
        "id": r[IDX["ID"]],
        "date_str": str(r[IDX["DATE"]]) if r[IDX["DATE"]] else "",
        "zhk": r[IDX["ZHK"]] or "",
        "tenure": r[IDX["TENURE"]] or "",
        "gender": r[IDX["GENDER"]] or "",
        "age": r[IDX["AGE"]] or "",
        "office_visit": r[IDX["OFFICE_VISIT"]] or "",
        "office_freq": r[IDX["OFFICE_FREQ"]] or "",
        "device": r[IDX["DEV"]] or "",
        "csat": _int(r[IDX["CSAT"]], 1, 5),
        "nps": _int(r[IDX["NPS"]], 0, 10),
        "opyt": _int(r[IDX["OPYT"]], 0, 10),
        "garant_ask": r[IDX["GARANT_ASK"]] or "",
        "garant_rate": _int(r[IDX["GARANT_RATE"]], 1, 5),
        "garant_impact": r[IDX["GARANT_IMPACT"]] or "",
        "open_ans": r[IDX["OPEN"]] or "",
        "csi": [_int(r[IDX["CSI_START"]+i*2], 1, 5) for i in range(13)],
    }
    # Парсим дату и час
    d = rec["date_str"]
    rec["date"] = None
    rec["hour"] = None
    if d:
        try:
            parts = d.split(" ")
            dp = parts[0]
            rec["date"] = dt.datetime.strptime(dp, "%d.%m.%Y").date()
            if len(parts) > 1:
                rec["hour"] = int(parts[1].split(":")[0])
        except: pass
    records.append(rec)

print(f"Записей: {len(records)}")

# --- Хелперы ---
def nps_calc(vals):
    vals = [v for v in vals if v is not None]
    n = len(vals)
    if n == 0: return None, 0, 0, 0, 0
    p = sum(1 for v in vals if v >= 9)
    d = sum(1 for v in vals if v <= 6)
    return round((p-d)/n*100, 1), n, p, n-p-d, d

def csat_calc(vals):
    vals = [v for v in vals if v is not None]
    n = len(vals)
    if n == 0: return None, 0
    return round(sum(1 for v in vals if v >= 4)/n*100, 1), n

def dist_0_10(vals):
    vals = [v for v in vals if v is not None]
    n = len(vals) or 1
    c = Counter(vals)
    return [c.get(s, 0) for s in range(11)], len(vals)

# --- Общий свод 2026 ---
nps_val, n_nps, p_cnt, pas_cnt, d_cnt = nps_calc([r["nps"] for r in records])
csat_val, n_csat = csat_calc([r["csat"] for r in records])
opyt_val, n_opyt, p_op, pas_op, d_op = nps_calc([r["opyt"] for r in records])
nps_dist, _ = dist_0_10([r["nps"] for r in records])
exp_dist, _ = dist_0_10([r["opyt"] for r in records])

# CSI средний
csi_all = []
for r in records:
    for v in r["csi"]:
        if v is not None: csi_all.append(v)
csi_score = round(sum(csi_all)/len(csi_all), 2) if csi_all else 0.0
print(f"NPS: {nps_val}, CSAT: {csat_val}, Индекс опыта: {opyt_val}, CSI: {csi_score}")
print(f"n_promo: {p_cnt} ({round(p_cnt/n_nps*100,1)}%), n_pas: {pas_cnt}, n_crit: {d_cnt}")

# --- Дневная динамика ---
by_date = defaultdict(list)
for r in records:
    if r["date"]: by_date[r["date"]].append(r)

daily = []
cumul_ans = 0
cumul_p = 0
cumul_d = 0
cumul_pas = 0
cumul_n = 0
prev_nps_cumul = None
prev_idx_cumul = None
sorted_dates = sorted(by_date.keys())
for i, d in enumerate(sorted_dates, 1):
    day_recs = by_date[d]
    n_day = len(day_recs)
    cumul_ans += n_day
    day_nps_vals = [r["nps"] for r in day_recs if r["nps"] is not None]
    day_p = sum(1 for v in day_nps_vals if v >= 9)
    day_d = sum(1 for v in day_nps_vals if v <= 6)
    day_pas = len(day_nps_vals) - day_p - day_d
    cumul_p += day_p
    cumul_d += day_d
    cumul_pas += day_pas
    cumul_n += len(day_nps_vals)
    nps_day = round((day_p - day_d)/len(day_nps_vals)*100) if day_nps_vals else 0
    nps_cumul = round((cumul_p - cumul_d)/cumul_n*100) if cumul_n else 0
    delta_nps = round(nps_cumul - prev_nps_cumul, 1) if prev_nps_cumul is not None else None
    idx_vals = [r["opyt"] for r in records if r["date"] and r["date"] <= d and r["opyt"] is not None]
    idx_p = sum(1 for v in idx_vals if v >= 9)
    idx_d = sum(1 for v in idx_vals if v <= 6)
    idx_cumul = round((idx_p - idx_d)/len(idx_vals)*100, 1) if idx_vals else 0
    delta_idx = round(idx_cumul - prev_idx_cumul, 1) if prev_idx_cumul is not None else None
    csat_vals = [r["csat"] for r in records if r["date"] and r["date"] <= d and r["csat"] is not None]
    csat_cumul = round(sum(1 for v in csat_vals if v >= 4)/len(csat_vals)*100, 1) if csat_vals else 0
    daily.append({
        "date": d.strftime("%Y-%m-%d"),
        "date_ru": d.strftime("%d.%m.%Y"),
        "day_num": i,
        "ans_day": n_day,
        "ans_cumul": cumul_ans,
        "nps": nps_cumul, "nps_cumul": nps_cumul, "nps_day": nps_day,
        "delta_nps": delta_nps,
        "day_p": day_p, "day_pas": day_pas, "day_d": day_d,
        "cumul_p": cumul_p, "cumul_pas": cumul_pas, "cumul_d": cumul_d,
        "cumul_p_pct": round(cumul_p/cumul_n*100, 1) if cumul_n else 0,
        "cumul_pas_pct": round(cumul_pas/cumul_n*100, 1) if cumul_n else 0,
        "cumul_d_pct": round(cumul_d/cumul_n*100, 1) if cumul_n else 0,
        "idx_cumul": idx_cumul, "delta_idx": delta_idx,
        "csat_cumul": csat_cumul,
    })
    prev_nps_cumul = nps_cumul
    prev_idx_cumul = idx_cumul

print(f"Дней: {len(daily)}")
for d in daily: print(f"  {d['date_ru']}: {d['ans_day']:4d} ответов (накоп {d['ans_cumul']}), NPS {d['nps_cumul']}")

# --- Активные часы ---
hours_data = []
hours_cnt = Counter()
for r in records:
    if r["hour"] is not None: hours_cnt[r["hour"]] += 1
total_h = sum(hours_cnt.values()) or 1
for h in range(24):
    cnt = hours_cnt.get(h, 0)
    hours_data.append({"hour": h, "n": cnt, "pct": round(cnt/total_h*100, 1)})
top_hours = sorted(hours_data, key=lambda x: -x["n"])[:5]
print(f"ТОП-5 часов: {[(h['hour'], h['n']) for h in top_hours]}")

# --- ЖК детально ---
zhk_agg = defaultdict(list)
for r in records:
    if r["zhk"]: zhk_agg[r["zhk"]].append(r)

zhk_rows = []
for zhk in sorted(zhk_agg.keys()):
    recs = zhk_agg[zhk]
    nps_v, n, p, pas, d = nps_calc([r["nps"] for r in recs])
    cs, _ = csat_calc([r["csat"] for r in recs])
    op, _, _, _, _ = nps_calc([r["opyt"] for r in recs])
    n2025 = zhk_2025.get(zhk, {}).get("n_2025", 0)
    nps2025 = zhk_2025.get(zhk, {}).get("nps_2025")
    delta = round(nps_v - nps2025, 1) if nps_v is not None and nps2025 is not None else None
    zhk_rows.append({
        "zhk": zhk, "n_2026": n, "nps_2026": nps_v, "csat_2026": cs, "opyt_2026": op,
        "n_2025": n2025, "nps_2025": nps2025, "delta": delta,
        "in_perim": zhk in PERIM_SET,
    })
n_stab = sum(1 for r in zhk_rows if r["n_2026"] >= 50 and r["in_perim"])
print(f"Стабилизировано: {n_stab} из 96")

# --- BCG-матрица ---
map_data = []
seg_cnt = Counter()
for i, p in enumerate(perimeter, 1):
    zhk = p["zhk"]
    zdata = next((r for r in zhk_rows if r["zhk"] == zhk), None)
    n = zdata["n_2026"] if zdata else 0
    nps26 = zdata["nps_2026"] if zdata else None
    n25 = zdata["n_2025"] if zdata else 0
    nps25 = zdata["nps_2025"] if zdata else None
    delta = zdata["delta"] if zdata else None
    n_proj = round(n * 3.0)
    risk = "нет" if n_proj >= 60 else ("средний" if n_proj >= 40 else "высокий")
    if n < 50:
        seg = "Недобор"
    elif nps26 is None:
        seg = "Недобор"
    elif nps26 >= 0:
        seg = "Звезды"
    elif delta is not None and delta >= 5:
        seg = "Восходящие"
    elif delta is not None and delta <= -5:
        seg = "Кризис"
    else:
        seg = "Спящие"
    seg_cnt[seg] += 1
    map_data.append({
        "num": i, "zhk": zhk, "region": p["reg"], "group": p["meropr"],
        "kpi": "", "n_2026": n, "nps_2026": nps26,
        "n_2025": n25, "nps_2025": nps25, "delta": delta,
        "seg": seg, "n_proj": n_proj, "risk": risk,
    })
print(f"Сегменты BCG: {dict(seg_cnt)}")

# --- Прогноз объёма ---
forecast_vol = []
for m in map_data:
    forecast_vol.append({
        "zhk": m["zhk"], "n_now": m["n_2026"], "n_proj": m["n_proj"],
        "risk": m["risk"], "region": m["region"],
    })
risk_cnt = Counter(m["risk"] for m in map_data)
print(f"Риск недобора: {dict(risk_cnt)}")

# --- Прогноз NPS к 30.09 ---
# Гранулярно: каждый ЖК свою кривую сохраняет. n_проекция * средний_NPS_ЖК
proj_p = 0
proj_d = 0
proj_pas = 0
for m in map_data:
    n_proj = m["n_proj"]
    nps26 = m["nps_2026"]
    if n_proj == 0 or nps26 is None:
        # Используем средний NPS периметра для незаполненных
        avg_p_rate = p_cnt / n_nps if n_nps else 0
        avg_d_rate = d_cnt / n_nps if n_nps else 0
        proj_p += n_proj * avg_p_rate
        proj_d += n_proj * avg_d_rate
        proj_pas += n_proj * (1 - avg_p_rate - avg_d_rate)
    else:
        # Раскладываем NPS на p/d при известной пропорции у этого ЖК
        recs = zhk_agg.get(m["zhk"], [])
        nps_vals = [r["nps"] for r in recs if r["nps"] is not None]
        n_r = len(nps_vals) or 1
        p_r = sum(1 for v in nps_vals if v >= 9) / n_r
        d_r = sum(1 for v in nps_vals if v <= 6) / n_r
        proj_p += n_proj * p_r
        proj_d += n_proj * d_r
        proj_pas += n_proj * (1 - p_r - d_r)

total_proj = proj_p + proj_d + proj_pas
nps_forecast = round((proj_p - proj_d)/total_proj*100, 1) if total_proj else 0
print(f"Прогноз NPS к 30.09: {nps_forecast} (проекция объёма {round(total_proj)})")

# --- Разрезы ---
def cut(field):
    grp = defaultdict(list)
    for r in records:
        k = r.get(field) or "—"
        grp[k].append(r)
    out = []
    for k, recs in grp.items():
        v, n, p, pas, d = nps_calc([r["nps"] for r in recs])
        out.append({"seg": k, "n": n, "nps": v, "p": p, "pas": pas, "d": d})
    return sorted(out, key=lambda x: -(x["n"] or 0))

cuts = {
    "gender": cut("gender"),
    "age": cut("age"),
    "tenure": cut("tenure"),
    "office_visit": cut("office_visit"),
    "office_freq": cut("office_freq"),
    "device": cut("device"),
}
print("\nРазрезы (пол):", cuts["gender"][:3])

# --- Факторы ---
promo = [r for r in records if r["nps"] is not None and r["nps"] >= 9]
crit  = [r for r in records if r["nps"] is not None and r["nps"] <= 6]
factors = []
for i, label in enumerate(CSI_LABELS):
    pv = [r["csi"][i] for r in promo if r["csi"][i] is not None]
    cv = [r["csi"][i] for r in crit  if r["csi"][i] is not None]
    all_v = [r["csi"][i] for r in records if r["csi"][i] is not None]
    p_avg = round(sum(pv)/len(pv), 2) if pv else None
    c_avg = round(sum(cv)/len(cv), 2) if cv else None
    a_avg = round(sum(all_v)/len(all_v), 2) if all_v else None
    gap = round(p_avg - c_avg, 2) if p_avg is not None and c_avg is not None else None
    factors.append({"name": label, "avg_all": a_avg, "avg_promo": p_avg, "avg_crit": c_avg, "gap": gap})
factors.sort(key=lambda x: -(x["gap"] or 0))
print(f"ТОП-3 факторы (макс gap промо−крит):")
for f in factors[:3]: print(f"  {f['name']}: promo={f['avg_promo']}, crit={f['avg_crit']}, gap={f['gap']}")

# --- CSI пульс ---
csi_pulse = []
for i, label in enumerate(CSI_LABELS):
    vals = [r["csi"][i] for r in records if r["csi"][i] is not None]
    avg = round(sum(vals)/len(vals), 2) if vals else None
    pct_max = round(avg/5*100, 1) if avg is not None else None
    csi_pulse.append({"name": label, "avg": avg, "pct_max": pct_max, "n": len(vals)})
csi_pulse.sort(key=lambda x: (x["avg"] or 0))
print("Худшие 3 критерия CSI:")
for c in csi_pulse[:3]: print(f"  {c['name']}: avg={c['avg']}, %max={c['pct_max']}")

# ============================================================
# ФОРМИРУЮ ВСЕ ОБЪЕКТЫ ДЛЯ СНИМКА
# ============================================================
now_ru = dt.datetime.now(dt.timezone(dt.timedelta(hours=3))).strftime("%d.%m.%Y %H:%M MSK")

SUMMARY_new = {
    "updated_at": now_ru,
    "total_responses": len(records),
    "nps_pct": nps_val,
    "index_pct": opyt_val,
    "csat_pct": csat_val,
    "csi_score": csi_score,
    "daily_2026": [
        {"date": d["date"], "day_num": d["day_num"],
         "ans_day": d["ans_day"], "ans_cumul": d["ans_cumul"],
         "nps": d["nps_cumul"], "nps_cumul": d["nps_cumul"], "nps_day": d["nps_day"]}
        for d in daily
    ],
    "n_nps": n_nps,
    "promoters_pct": round(p_cnt/n_nps*100, 1) if n_nps else 0,
    "neutrals_pct": round(pas_cnt/n_nps*100, 1) if n_nps else 0,
    "critics_pct": round(d_cnt/n_nps*100, 1) if n_nps else 0,
    "nps_dist": nps_dist,
    "n_exp": n_opyt,
    "exp_dist": exp_dist,
    "nps_forecast": nps_forecast,
    "hours": hours_data,
    "top_hours": top_hours,
    "stab_n": n_stab,
    "stab_total": 96,
}

# ============================================================
# ОБНОВЛЯЮ SHEETS
# ============================================================
html = open(SNAP, encoding="utf-8").read()

# Читаю текущий SHEETS
m = re.search(r'const SHEETS = (\{.*?\});\n', html, re.DOTALL)
if not m:
    raise SystemExit("SHEETS не найден")
SHEETS = json.loads(m.group(1))

# --- ПУЛЬС_ДЕНЬ ---
def fmt_pct(v):
    if v is None: return ""
    return f"{v:.1f}%".replace(".",",").replace("-","−")

pulse_header = SHEETS["Пульс_день"][0:3]  # заголовок + подзаголовок + шапка колонок
# Комментарии из старой версии — сохраняю
old_pulse = SHEETS["Пульс_день"][3:]
old_comments = {}
for row in old_pulse:
    if isinstance(row, list) and len(row) > 25 and row[0]:
        old_comments[str(row[0])] = row[25]

new_pulse = list(pulse_header)
for d in daily:
    day_recs = by_date[dt.datetime.strptime(d["date"], "%Y-%m-%d").date()]
    day_nps_vals = [r["nps"] for r in day_recs if r["nps"] is not None]
    day_csat_vals = [r["csat"] for r in day_recs if r["csat"] is not None]
    csat_day = round(sum(1 for v in day_csat_vals if v>=4)/len(day_csat_vals)*100,1) if day_csat_vals else 0
    dist = Counter(day_nps_vals)
    n_day_valid = len(day_nps_vals) or 1
    row = [
        d["date_ru"],
        d["ans_cumul"],
        d["ans_day"],
        fmt_pct(d["nps_cumul"]),
        fmt_pct(d["delta_nps"]) if d["delta_nps"] is not None else "",
        d["cumul_d"], d["cumul_pas"], d["cumul_p"],
        fmt_pct(d["cumul_d_pct"]),
        fmt_pct(d["cumul_pas_pct"]),
        fmt_pct(d["cumul_p_pct"]),
        fmt_pct(d["idx_cumul"]),
        fmt_pct(d["delta_idx"]) if d["delta_idx"] is not None else "",
        fmt_pct(d["csat_cumul"]),
    ]
    # 0..10 дневное распределение
    for s in range(11):
        cnt = dist.get(s, 0)
        row.append(fmt_pct(round(cnt/n_day_valid*100, 1)))
    # Коммуникация — сохраняю старые
    comm = old_comments.get(d["date_ru"], "")
    row.append(comm)
    new_pulse.append(row)

# Добавляю блок «Активные часы»
new_pulse.append([""])
new_pulse.append(["АКТИВНЫЕ ЧАСЫ (когда пользователи проходят опрос)"])
new_pulse.append(["Час","Ответов","% от всех"] + [""]*23)
for h in hours_data:
    marker = " ⚡ ПИК" if h in top_hours else ""
    new_pulse.append([f"{h['hour']:02d}:00", h["n"], fmt_pct(h["pct"]) + marker] + [""]*23)

SHEETS["Пульс_день"] = new_pulse

# --- Разрезы ---
razrez_header = SHEETS["Разрезы"][:3] if len(SHEETS["Разрезы"])>=3 else [["РАЗРЕЗЫ NPS-2026"], [""], ["Сегмент","n","NPS","Промо","Нейтр","Крит"]]
razrez_new = list(razrez_header)
for cut_name, cut_label in [("gender","ПОЛ"),("age","ВОЗРАСТ"),("tenure","СРОК ВЛАДЕНИЯ"),("office_visit","ВИЗИТ В ОФИС"),("office_freq","ЧАСТОТА ВИЗИТОВ"),("device","УСТРОЙСТВО")]:
    razrez_new.append([""])
    razrez_new.append([cut_label])
    for c in cuts[cut_name]:
        hyp = " ⚠" if c["n"] < 50 else ""
        razrez_new.append([str(c["seg"])+hyp, c["n"], fmt_pct(c["nps"]) if c["nps"] is not None else "—", c["p"], c["pas"], c["d"]])
SHEETS["Разрезы"] = razrez_new

# --- Факторы ---
fact_new = [
    ["ФАКТОРЫ NPS — 13 КРИТЕРИЕВ (вопрос 8)"],
    ["Сортировано по гэпу «промо − крит» (что сильнее всего расщепляет NPS)"],
    ["Критерий","Средняя (все)","Средняя (промо)","Средняя (крит)","Гэп"],
]
for f in factors:
    fact_new.append([f["name"], f["avg_all"] or "—", f["avg_promo"] or "—", f["avg_crit"] or "—", f["gap"] or "—"])
SHEETS["Факторы"] = fact_new

# --- ЖК_детально ---
zhk_new = [
    ["ЖК ДЕТАЛЬНО — все ЖК (n<50 маркированы ⚠ гипотеза)"],
    ["Отсортировано по количеству ответов"],
    ["ЖК","n 2026","NPS 2026","CSAT %","Индекс опыта","n 2025","NPS 2025","Δ YoY"],
]
for r in sorted(zhk_rows, key=lambda x: -x["n_2026"]):
    hyp = " ⚠" if r["n_2026"] < 50 else ""
    zhk_new.append([
        r["zhk"]+hyp,
        r["n_2026"],
        fmt_pct(r["nps_2026"]) if r["nps_2026"] is not None else "—",
        fmt_pct(r["csat_2026"]) if r["csat_2026"] is not None else "—",
        fmt_pct(r["opyt_2026"]) if r["opyt_2026"] is not None else "—",
        r["n_2025"],
        fmt_pct(r["nps_2025"]) if r["nps_2025"] is not None else "—",
        fmt_pct(r["delta"]) if r["delta"] is not None else "—",
    ])
SHEETS["ЖК_детально"] = zhk_new

# --- Индекс_опыта ---
idx_new = [
    ["ИНДЕКС ОПЫТА (вопрос 13) — накопительно 2026"],
    [f"Значение: {opyt_val}% · n={n_opyt}"],
    ["Оценка","Ответов","%"],
]
for s in range(11):
    cnt = exp_dist[s]
    pct = round(cnt/n_opyt*100, 1) if n_opyt else 0
    idx_new.append([s, cnt, fmt_pct(pct)])
SHEETS["Индекс_опыта"] = idx_new

# --- CSI_пульс ---
csi_new = [
    ["CSI ПУЛЬС — 13 критериев вопроса 8"],
    [f"Средний CSI: {csi_score} из 5.00"],
    ["Критерий","Средняя","% от макс","n ответов"],
]
for c in csi_pulse:
    csi_new.append([c["name"], c["avg"] or "—", fmt_pct(c["pct_max"]) if c["pct_max"] is not None else "—", c["n"]])
SHEETS["CSI_пульс"] = csi_new

# --- Прогноз_NPS ---
prog_new = [
    ["ПРОГНОЗ NPS К 30.09.2026"],
    [f"Гранулярная модель по 96 ЖК. KPI 2025 = −32 п.п. (красная линия) · Цель 2026 ≤ −20 п.п."],
    ["Показатель","Значение"],
    ["NPS сейчас (накоп.)", fmt_pct(nps_val)],
    ["Прогноз NPS к 30.09 (модель)", fmt_pct(nps_forecast)],
    ["Проекция объёма ответов к 30.09", int(round(total_proj))],
    ["Гэп до −32 (KPI 2025)", fmt_pct(round(nps_forecast - (-32), 1))],
    ["Гэп до −35 (запас безопасности)", fmt_pct(round(nps_forecast - (-35), 1))],
    ["Гэп до −20 (цель 2026)", fmt_pct(round(nps_forecast - (-20), 1))],
    [""],
    ["СЦЕНАРИИ ПО ЦЕЛЯМ"],
    ["Цель NPS","Нужно ещё ответов","Средний NPS дозаборки"],
]
n_remaining = int(round(total_proj)) - n_nps
for target in [-35, -32, -30, -25, -20]:
    # средний NPS дозаборки при котором финал будет == target
    if n_remaining > 0:
        # (nps_val * n_nps + nps_avg_new * n_remaining) / (n_nps + n_remaining) = target
        # nps_avg_new = (target*(n_nps+n_remaining) - nps_val*n_nps) / n_remaining
        nps_avg_new = round((target*(n_nps+n_remaining) - nps_val*n_nps)/n_remaining, 1)
    else:
        nps_avg_new = None
    safe_plan = int(round(n_remaining * 1.05)) if n_remaining > 0 else 0
    prog_new.append([f"NPS = {target}", safe_plan, fmt_pct(nps_avg_new) if nps_avg_new is not None else "—"])

# Блок прогноза объёма по ЖК
prog_new.append([""])
prog_new.append(["ПРОГНОЗ ОБЪЁМА И РИСК НЕДОБОРА (n_проекция = n_текущий × 3.0)"])
prog_new.append([f"Высокий риск: {risk_cnt.get('высокий', 0)} · Средний: {risk_cnt.get('средний', 0)} · Безопасно: {risk_cnt.get('нет', 0)}"])
prog_new.append(["ЖК","n сейчас","n проекция к 30.09","Риск"])
for m in sorted(map_data, key=lambda x: x["n_proj"]):
    prog_new.append([m["zhk"], m["n_2026"], m["n_proj"], m["risk"]])
SHEETS["Прогноз_NPS"] = prog_new

# --- Списки_для_работы ---
lists_new = [
    ["СПИСКИ ДЛЯ РАБОТЫ — открытые ответы"],
    ["Отфильтровано: не пустые"],
    ["ЖК","Сегмент","NPS","Комментарий (вопрос 14)"],
]
for r in records:
    txt = str(r["open_ans"]).strip() if r["open_ans"] else ""
    if not txt or len(txt) < 3: continue
    if r["nps"] is None: continue
    seg = "промо" if r["nps"] >= 9 else ("нейтрал" if r["nps"] >= 7 else "критик")
    lists_new.append([r["zhk"], seg, r["nps"], txt[:400]])
SHEETS["Списки_для_работы"] = lists_new

# --- Недельный_синтез ---
weekly = defaultdict(list)
for r in records:
    if r["date"]:
        iso = r["date"].isocalendar()
        weekly[f"W{iso.week}"].append(r)
weekly_new = [
    ["НЕДЕЛЬНЫЙ СИНТЕЗ"],
    ["Свёртка Пульса по неделям"],
    ["Неделя","Ответов","NPS","CSAT %","Индекс опыта"],
]
for wk in sorted(weekly.keys()):
    recs = weekly[wk]
    v,n,_,_,_ = nps_calc([r["nps"] for r in recs])
    cs,_ = csat_calc([r["csat"] for r in recs])
    op,_,_,_,_ = nps_calc([r["opyt"] for r in recs])
    weekly_new.append([wk, n, fmt_pct(v) if v is not None else "—", fmt_pct(cs) if cs is not None else "—", fmt_pct(op) if op is not None else "—"])
SHEETS["Недельный_синтез"] = weekly_new

# --- Справочник_ЖК ---
ref_new = [
    ["СПРАВОЧНИК ЖК — 96 в периметре"],
    [""],
    ["№","ЖК","Регион","Программа"],
]
for i, p in enumerate(perimeter, 1):
    ref_new.append([i, p["zhk"], p["reg"], p["meropr"]])
SHEETS["Справочник_ЖК"] = ref_new

# --- Методика — оставляю как есть если есть, иначе минимальную ---
if "Методика" not in SHEETS or len(SHEETS["Методика"]) < 3:
    SHEETS["Методика"] = [
        ["МЕТОДИКА РАСЧЁТА"],
        ["Все показатели пересчитываются из первичных ответов Testograf"],
        ["Показатель","Формула"],
        ["NPS", "%(9-10) − %(0-6) по вопросу 9"],
        ["CSAT", "% оценок 4-5 по вопросу 7 (шкала 1-5)"],
        ["Индекс опыта", "%(9-10) − %(0-6) по вопросу 13"],
        ["CSI средний", "Среднее по 13 подкритериям вопроса 8"],
        ["Стабилизация ЖК", "n ≥ 50 респондентов"],
        ["Сегментация BCG", "Звёзды NPS≥0, Восходящие ΔYoY≥+5, Кризис Δ≤−5, Спящие иное, Недобор n<50"],
        ["Прогноз объёма", "n_проекция = n_текущий × 3.0 (осталось ~20 дней)"],
        ["KPI 2025 факт", "−32 п.п. (красная линия, нельзя ниже)"],
    ]

# ============================================================
# ВСТАВЛЯЮ ОБРАТНО В СНИМОК
# ============================================================
# Заменяю SUMMARY
new_summary_json = json.dumps(SUMMARY_new, ensure_ascii=False)
html = re.sub(r'const SUMMARY = \{.*?\};\n', f'const SUMMARY = {new_summary_json};\n', html, count=1, flags=re.DOTALL)

# Очищаю управляющие символы в строках (иначе JSON.parse не пройдёт)
def _clean(v):
    if isinstance(v, str):
        return v.replace("\r"," ").replace("\n"," ").replace("\t"," ")
    elif isinstance(v, list):
        return [_clean(x) for x in v]
    elif isinstance(v, dict):
        return {k: _clean(vv) for k,vv in v.items()}
    return v
SHEETS = _clean(SHEETS)

# Заменяю SHEETS
new_sheets_json = json.dumps(SHEETS, ensure_ascii=False)
html = re.sub(r'const SHEETS = \{.*?\};\n', f'const SHEETS = {new_sheets_json};\n', html, count=1, flags=re.DOTALL)

# Заменяю MAP_ZHK
new_map_json = json.dumps(map_data, ensure_ascii=False)
html = re.sub(r'const MAP_ZHK = \[.*?\];\n', f'const MAP_ZHK = {new_map_json};\n', html, count=1, flags=re.DOTALL)

# Заменяю FORECAST_VOL
new_fv_json = json.dumps(forecast_vol, ensure_ascii=False)
html = re.sub(r'const FORECAST_VOL = \[.*?\];\n', f'const FORECAST_VOL = {new_fv_json};\n', html, count=1, flags=re.DOTALL)

# Заменяю STAB_TOTAL
html = re.sub(r'const STAB_TOTAL = \{.*?\};\n', f'const STAB_TOTAL = {{"n": {n_stab}, "total": 96}};\n', html, count=1, flags=re.DOTALL)

# Также обновляю json файл сегментов
with open("/home/user/workspace/nps_pipeline/_snapshots/map_zhk.json", "w", encoding="utf-8") as f:
    json.dump(map_data, f, ensure_ascii=False, indent=2)

open(SNAP, "w", encoding="utf-8").write(html)
print(f"\n✓ Снимок обновлён: {SNAP}")
print(f"✓ Обновлено: {now_ru}")
print(f"✓ Всего ответов: {len(records)}")
print(f"✓ NPS: {nps_val}, Прогноз: {nps_forecast}")
print(f"✓ Стабилизировано: {n_stab}/96")